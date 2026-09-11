"""
Genera el checklist de controles de seguridad del módulo 2 en dos formatos
desde una única fuente de datos (la lista CONTROLES de abajo):

  - Checklist_Seguridad.xlsx  -> se proyecta y se llena EN VIVO en la sesión.
                                         Clasifica cada respuesta automáticamente.
  - CHECKLIST.md                      -> versión legible en GitHub.

Para agregar o ajustar un control: editar CONTROLES y volver a correr
    python generar_checklist.py
Luego recalcular el Excel (LibreOffice o abrirlo y guardarlo en Excel).

Regla de clasificación (la misma que se explica en la sala):
    Sí               -> Cumple
    No sé            -> Riesgo a validar
    No / Parcial     -> Quick win       si el esfuerzo es Bajo
                        Profundización  si el esfuerzo es Medio o Alto
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).parent

NIVEL = {"O": "Obligatorio", "R": "Recomendado", "Op": "Opcional", "NA": "N/A"}

DOMINIOS = [
    ("R", "Red y segmentación"),
    ("B", "Borde: balanceador y WAF"),
    ("I", "Identidad y segregación de funciones"),
    ("S", "Secretos y cifrado"),
    ("A", "Acceso administrativo"),
    ("D", "Seguridad de datos"),
    ("P", "Postura"),
    ("L", "Registro y respuesta"),
    ("M", "Operación para terceros"),
]

# (id, control, por qué importa, dev, qa, prod, esfuerzo, cómo se verifica, núcleo)
CONTROLES = [
    # --- Red y segmentación ---
    ("R-01", "Capas de aplicación y datos en subredes privadas; solo el balanceador tiene IP pública",
     "Una IP pública en un servidor de aplicación es superficie de ataque que no aporta nada: el tráfico legítimo entra por el balanceador.",
     "R", "O", "O", "Medio", "Script de auditoría (R-01) · subredes con prohibit-public-ip-on-vnic = true", True),
    ("R-02", "Security list por defecto de cada VCN sin SSH/RDP abierto a 0.0.0.0/0",
     "OCI crea cada VCN con una security list por defecto que abre el 22 a internet. Una subred creada sin especificar lista la hereda.",
     "O", "O", "O", "Bajo", "Script de auditoría (R-02)", True),
    ("R-03", "Reglas de tráfico por NSG según función (balanceador → app → datos), no listas amplias por subred",
     "Con NSG, «solo el balanceador habla con la app» se escribe literalmente; con listas por subred se termina abriendo la subred entera.",
     "R", "O", "O", "Medio", "Script de auditoría (R-03) · VCN → Network Security Groups", True),
    ("R-04", "Salida a internet solo por NAT, a servicios OCI por Service Gateway; la capa de datos sin salida a internet",
     "Un servidor comprometido lo primero que intenta es salir. Controlar la salida limita el daño y la exfiltración.",
     "Op", "R", "O", "Bajo", "Tablas de ruteo de subredes privadas: sin Internet Gateway", False),
    ("R-05", "Ambientes separados por compartment y VCN, sin conectividad abierta entre desarrollo y producción",
     "Un error o una credencial filtrada en desarrollo no debe tener camino hacia producción.",
     "O", "O", "O", "Medio", "Compartments por ambiente · revisar peering y DRG entre VCNs", True),

    # --- Borde ---
    ("B-01", "Solo el balanceador (o un API Gateway) publica servicios a internet",
     "Un único punto de entrada es un único lugar donde aplicar TLS, WAF, límites y registro.",
     "O", "O", "O", "Medio", "Inventario de IPs públicas: oci network public-ip list", False),
    ("B-02", "WAF con reglas OWASP delante de toda aplicación pública",
     "Filtra ataques comunes (inyección, XSS) antes de que lleguen al código y compra tiempo cuando aparece una vulnerabilidad nueva.",
     "Op", "R", "O", "Bajo", "Script de auditoría (B-02)", True),
    ("B-03", "TLS terminado en el balanceador con certificado gestionado; HTTP solo redirige a HTTPS",
     "Credenciales y sesiones viajan cifradas. Los certificados gestionados evitan caducidades sorpresa.",
     "R", "O", "O", "Bajo", "Script de auditoría (B-03) · listeners del balanceador", True),
    ("B-04", "Límite de tasa en endpoints sensibles (inicio de sesión, APIs públicas)",
     "Frena fuerza bruta y abuso de APIs sin tocar la aplicación.",
     "Op", "R", "R", "Bajo", "Política WAF: reglas de rate limiting", False),

    # --- Identidad ---
    ("I-01", "MFA obligatorio para todo acceso a la consola",
     "La mayoría de incidentes en nube empiezan con una credencial robada. MFA la vuelve insuficiente por sí sola.",
     "O", "O", "O", "Bajo", "Script de auditoría (I-01) · dominio de identidad → política de inicio de sesión", True),
    ("I-02", "Nadie opera el día a día con el grupo Administrators; cuentas de emergencia documentadas y vigiladas",
     "El administrador total es la llave maestra: se usa en emergencias, y cuando se usa, alguien se entera.",
     "O", "O", "O", "Bajo", "Miembros del grupo Administrators: pocos, nominales y justificados", True),
    ("I-03", "Grupos por función (red, cómputo, datos, seguridad, auditoría) con políticas limitadas a su compartment",
     "Segregación de funciones: quien despliega aplicaciones no cambia reglas de red ni permisos.",
     "R", "O", "O", "Medio", "Políticas con «manage all-resources in tenancy» fuera de Administrators", True),
    ("I-04", "La automatización (Terraform, CI/CD, scripts) usa principals de instancia o recurso, no llaves API de personas",
     "Una llave API personal en un pipeline es una credencial sin dueño el día que esa persona se va.",
     "R", "O", "O", "Medio", "Llaves API en variables de CI y archivos ~/.oci de servidores", True),
    ("I-05", "Llaves API rotadas cada 90 días como máximo; usuarios inactivos revisados cada trimestre",
     "Las credenciales viejas son las que nadie recuerda tener.",
     "R", "R", "O", "Bajo", "Script de auditoría (I-05)", False),

    # --- Secretos ---
    ("S-01", "Ningún secreto en código, tfvars, cloud-init ni variables de CI en texto plano; se consumen desde OCI Vault",
     "El repositorio es el primer lugar donde busca un atacante, y el historial de Git no olvida.",
     "O", "O", "O", "Medio", 'git log -p | grep -iE "password|secret|BEGIN .*KEY"', True),
    ("S-02", "Estado de Terraform almacenado remoto, cifrado y con acceso restringido",
     "El tfstate guarda en claro contraseñas y llaves de lo que crea. Suele terminar en un portátil o en un bucket abierto.",
     "R", "O", "O", "Bajo", "¿Dónde vive el tfstate? Backend en Object Storage con acceso por política", True),
    ("S-03", "Rotación definida de secretos (credenciales de base de datos, tokens) con responsable",
     "Un secreto que nunca rota termina filtrándose sin que nadie lo note.",
     "Op", "R", "R", "Medio", "Vault: fecha de la última versión de cada secreto", False),
    ("S-04", "Llaves propias en Vault para datos de producción cuando lo exija un contrato o una regulación",
     "OCI cifra todo por defecto; la llave propia agrega control de revocación y evidencia para auditorías.",
     "Op", "Op", "R", "Medio", "Volúmenes, buckets y bases: campo kms-key-id", False),

    # --- Acceso administrativo ---
    ("A-01", "Acceso administrativo por OCI Bastion con sesiones que expiran; sin SSH/RDP expuesto",
     "El 22 abierto no es una necesidad operativa, es una costumbre. Bastion da el mismo acceso con identidad, TTL y registro.",
     "R", "O", "O", "Bajo", "Script de auditoría (R-02, R-03) · Identity & Security → Bastion", True),
    ("A-02", "Llaves SSH por persona, nunca compartidas; revocación inmediata cuando alguien sale del equipo",
     "Una llave compartida hace imposible saber quién hizo qué y obliga a rotar para todos.",
     "O", "O", "O", "Bajo", "authorized_keys en los servidores: ¿una llave por persona?", True),
    ("A-03", "Parcheo de sistema operativo con cadencia definida y medida",
     "La mayoría de explotaciones usan vulnerabilidades con parche disponible hace meses.",
     "R", "O", "O", "Medio", "OS Management Hub o equivalente: % de hosts al día", False),

    # --- Datos ---
    ("D-01", "Bases de datos sin endpoint público; acceso solo desde el NSG de aplicación",
     "Una base expuesta a internet recibe intentos de acceso en minutos.",
     "O", "O", "O", "Bajo", "Consola de la base: sin IP pública · NSG con el puerto solo desde la app", True),
    ("D-02", "Buckets de Object Storage sin acceso público, salvo contenido intencionalmente público y documentado",
     "Es la fuga de datos más común y más evitable de la nube.",
     "O", "O", "O", "Bajo", "Script de auditoría (D-02)", True),
    ("D-03", "La aplicación usa un usuario de base de datos con mínimo privilegio, nunca el administrador",
     "Si la aplicación se compromete, el atacante hereda exactamente los permisos de ese usuario.",
     "R", "O", "O", "Medio", "Usuarios y privilegios definidos en la base", True),
    ("D-04", "Datos de producción no se copian a desarrollo o pruebas sin enmascaramiento",
     "Desarrollo tiene casi siempre controles más débiles. Datos reales ahí son producción sin protección.",
     "O", "O", "NA", "Medio", "Procedimiento de refresco de ambientes", True),
    ("D-05", "Evaluación periódica de configuración y usuarios de la base (Data Safe o equivalente)",
     "Detecta usuarios sobrantes, privilegios excesivos y configuraciones débiles antes que un auditor.",
     "Op", "R", "R", "Bajo", "Data Safe: evaluación de seguridad y de usuarios [VALIDAR soporte según motor]", False),
    ("D-06", "Respaldos cifrados, fuera del alcance de las credenciales de producción, con restauración probada cada trimestre",
     "Un respaldo que nunca se restauró es una hipótesis. Y un ransomware que alcanza los respaldos los cifra también.",
     "R", "R", "O", "Medio", "Última restauración de prueba documentada", False),

    # --- Postura ---
    ("P-01", "Cloud Guard habilitado con detectores de configuración y actividad sobre todos los compartments",
     "Detecta de forma continua lo que un checklist revisa una vez. No tiene costo adicional.",
     "O", "O", "O", "Bajo", "Script de auditoría (P-01) · Cloud Guard → Problems", True),
    ("P-02", "Security Zones en los compartments de producción",
     "Pasa de detectar a prevenir: la plataforma rechaza lo que viola la política antes de que exista.",
     "Op", "R", "O", "Medio", "Security Zones asociadas a los compartments de producción", True),
    ("P-03", "Escaneo de vulnerabilidades de hosts e imágenes de contenedor",
     "Sin inventario de vulnerabilidades no hay forma de priorizar parches.",
     "R", "R", "R", "Bajo", "Vulnerability Scanning: recetas y objetivos configurados", False),
    ("P-04", "Revisión periódica contra el CIS OCI Foundations Benchmark",
     "Da una línea base externa y reconocida para medir avance y responder auditorías.",
     "Op", "R", "R", "Medio", "Reporte de cumplimiento CIS del repositorio OCI CIS Landing Zone", False),

    # --- Registro y respuesta ---
    ("L-01", "Retención de Audit en 365 días",
     "Muchos incidentes se descubren meses después. Sin registro no hay investigación posible.",
     "O", "O", "O", "Bajo", "Script de auditoría (L-01) · Tenancy → Audit retention", True),
    ("L-02", "Flow logs de VCN en subredes de producción",
     "Permiten reconstruir quién habló con quién durante un incidente.",
     "Op", "R", "R", "Bajo", "Script de auditoría (L-02)", False),
    ("L-03", "Alertas ante cambios críticos: políticas IAM, security lists, NSGs y gateways",
     "Un cambio de regla a las 2 a. m. debería despertar a alguien, no descubrirse en la preparación.",
     "R", "O", "O", "Bajo", "Events + Notifications con reglas sobre esos tipos de evento", True),
    ("L-04", "Procedimiento de respuesta a incidentes con responsables, contactos y primeros pasos",
     "En un incidente, la primera hora se pierde decidiendo quién decide.",
     "R", "R", "O", "Medio", "Documento vigente y probado en un simulacro", True),
    ("L-05", "Registros centralizados y enviados a SIEM o Logging Analytics con retención definida",
     "Correlacionar Audit, red y aplicación es lo que convierte registros en detección.",
     "Op", "R", "R", "Medio", "Service Connector Hub hacia SIEM o Logging Analytics", False),

    # --- Operación para terceros ---
    ("M-01", "Cada cliente final en tenancy o compartment propio; nunca recursos de dos clientes en el mismo compartment",
     "Aísla impacto, permisos y costo por cliente. Es también la base del showback por cliente del módulo 1.",
     "R", "O", "O", "Medio", "Estructura de compartments o tenancies por cliente", True),
    ("M-02", "El equipo de la organización entra a los tenancies de clientes con identidades nominales y federadas, no usuarios compartidos",
     "Un usuario compartido entre operadores impide atribuir acciones y cumplirle al cliente.",
     "O", "O", "O", "Medio", "Usuarios en tenancies de clientes: ¿nominales? ¿federados?", True),
    ("M-03", "Registro revisable de qué operador accedió a qué cliente y cuándo",
     "Es lo primero que pide un cliente regulado (salud, sector cooperativo) antes de delegar su operación.",
     "R", "O", "O", "Bajo", "Audit por tenancy, con revisión periódica", False),
]

# --- Estilos -----------------------------------------------------------------

FUENTE = "Arial"
AZUL = "1F3864"
GRIS_TXT = "595959"
AMARILLO_INPUT = "FFF2CC"
FILL_HEAD = PatternFill("solid", start_color=AZUL)
FILL_DOM = PatternFill("solid", start_color="D9E1F2")
FILL_INPUT = PatternFill("solid", start_color=AMARILLO_INPUT)
FILL_NUCLEO = PatternFill("solid", start_color="E2EFDA")
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

COLORES_CLASIF = {
    "Cumple": ("D9D9D9", "404040"),
    "Quick win": ("C6EFCE", "006100"),
    "Riesgo a validar": ("FFEB9C", "7F6000"),
    "Profundización": ("BDD7EE", "1F3864"),
}


def f(bold=False, size=10, color="000000", italic=False):
    return Font(name=FUENTE, bold=bold, size=size, color=color, italic=italic)


def construir_xlsx(ruta):
    wb = Workbook()
    ws_leeme = wb.active
    ws_leeme.title = "Leeme"
    ws = wb.create_sheet("Checklist")
    wr = wb.create_sheet("Resultado")

    n = len(CONTROLES)
    PRIMERA = 5
    ULTIMA = PRIMERA + n - 1
    rango = lambda col: f"Checklist!${col}${PRIMERA}:${col}${ULTIMA}"

    # ============================ CHECKLIST ============================
    ws["A1"] = "Checklist de controles de seguridad — Taller de arquitectura en OCI · módulo 2"
    ws["A1"].font = f(True, 14, AZUL)
    ws["A2"] = ("Celdas amarillas: se llenan en la sesión. Filtrar «Núcleo = Sí» para la revisión de 40 minutos; "
                "el resto queda como autoevaluación. La columna Clasificación se calcula sola.")
    ws["A2"].font = f(size=9, color=GRIS_TXT, italic=True)

    encabezados = ["ID", "Dominio", "Control", "Por qué importa", "Dev", "QA", "Prod",
                   "Esfuerzo", "Cómo se verifica", "Núcleo", "Estado la organización",
                   "Clasificación", "Dueño", "Nota", "_qw", "_rv", "_pr"]
    anchos = [7, 20, 48, 48, 12, 12, 12, 10, 38, 8, 14, 18, 22, 30, 5, 5, 5]
    for i, (h, w) in enumerate(zip(encabezados, anchos), start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = f(True, 10, "FFFFFF")
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[4].height = 30

    nombre_dom = dict(DOMINIOS)
    for idx, (cid, ctrl, porque, dev, qa, prod, esf, verif, nucleo) in enumerate(CONTROLES):
        r = PRIMERA + idx
        valores = [cid, nombre_dom[cid[0]], ctrl, porque, NIVEL[dev], NIVEL[qa], NIVEL[prod],
                   esf, verif, "Sí" if nucleo else "No", None]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = f(bold=(col == 1))
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="center" if col in (1, 5, 6, 7, 8, 10, 11) else "left")
            c.border = BORDE
        if nucleo:
            ws.cell(row=r, column=10).fill = FILL_NUCLEO
        # Entradas
        for col in (11, 13, 14):
            c = ws.cell(row=r, column=col)
            c.fill = FILL_INPUT
            c.border = BORDE
            c.font = f()
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="center" if col == 11 else "left")
        # Clasificación
        c = ws.cell(row=r, column=12,
                    value=(f'=IF(K{r}="","Pendiente",IF(K{r}="Sí","Cumple",'
                           f'IF(K{r}="No sé","Riesgo a validar",'
                           f'IF(H{r}="Bajo","Quick win","Profundización"))))'))
        c.font = f(True)
        c.alignment = Alignment(horizontal="center", vertical="top")
        c.border = BORDE
        # Auxiliares para las listas de Resultado
        for col, etiqueta in ((15, "Quick win"), (16, "Riesgo a validar"), (17, "Profundización")):
            L = get_column_letter(col)
            ws.cell(row=r, column=col,
                    value=f'=IF($L{r}="{etiqueta}",COUNTIF($L${PRIMERA}:$L{r},"{etiqueta}"),"")').font = f(size=8, color="A6A6A6")
        ws.row_dimensions[r].height = 48

    # Columnas auxiliares ocultas
    for col in ("O", "P", "Q"):
        ws.column_dimensions[col].hidden = True
    # "Por qué" y "Cómo se verifica" agrupadas: se pueden colapsar al proyectar
    ws.column_dimensions.group("D", "D", hidden=False, outline_level=1)
    ws.column_dimensions.group("I", "I", hidden=False, outline_level=1)

    dv_estado = DataValidation(type="list", formula1='"Sí,Parcial,No,No sé"', allow_blank=True,
                               showErrorMessage=True, errorTitle="Valor no válido",
                               error="Usar: Sí, Parcial, No o No sé")
    dv_nivel = DataValidation(type="list", formula1='"Obligatorio,Recomendado,Opcional,N/A"', allow_blank=False)
    ws.add_data_validation(dv_estado)
    ws.add_data_validation(dv_nivel)
    dv_estado.add(f"K{PRIMERA}:K{ULTIMA}")
    dv_nivel.add(f"E{PRIMERA}:G{ULTIMA}")

    for etiqueta, (fondo, texto) in COLORES_CLASIF.items():
        ws.conditional_formatting.add(
            f"L{PRIMERA}:L{ULTIMA}",
            FormulaRule(formula=[f'$L{PRIMERA}="{etiqueta}"'],
                        fill=PatternFill("solid", start_color=fondo),
                        font=Font(name=FUENTE, bold=True, color=texto)))
    ws.conditional_formatting.add(
        f"E{PRIMERA}:G{ULTIMA}",
        FormulaRule(formula=[f'E{PRIMERA}="Obligatorio"'], font=Font(name=FUENTE, bold=True, color="C00000")))

    ws.auto_filter.ref = f"A4:N{ULTIMA}"
    ws.freeze_panes = "D5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"

    # ============================ RESULTADO ============================
    wr["A1"] = "Resultado del módulo 2 — se actualiza solo mientras se llena el checklist"
    wr["A1"].font = f(True, 14, AZUL)
    for col, w in zip("ABCDEFG", [22, 12, 12, 12, 14, 12, 12]):
        wr.column_dimensions[col].width = w

    # KPIs
    wr["A3"] = "Indicador"
    wr["B3"] = "Valor"
    for c in (wr["A3"], wr["B3"]):
        c.font = f(True, 10, "FFFFFF"); c.fill = FILL_HEAD; c.border = BORDE
    kpis = [
        ("Controles evaluados", f'=COUNTA({rango("K")})'),
        ("  de ellos, del núcleo", f'=COUNTIFS({rango("J")},"Sí",{rango("K")},"<>")'),
        ("Cumple", f'=COUNTIF({rango("L")},"Cumple")'),
        ("Quick win", f'=COUNTIF({rango("L")},"Quick win")'),
        ("Riesgo a validar", f'=COUNTIF({rango("L")},"Riesgo a validar")'),
        ("Profundización", f'=COUNTIF({rango("L")},"Profundización")'),
        ("Pendiente de evaluar", f'=COUNTIF({rango("L")},"Pendiente")'),
        ("Brechas en obligatorios de producción",
         f'=COUNTIFS({rango("G")},"Obligatorio",{rango("K")},"<>",{rango("L")},"<>Cumple")'),
    ]
    for i, (et, fo) in enumerate(kpis, start=4):
        wr.cell(row=i, column=1, value=et).font = f(bold=(i == 11))
        c = wr.cell(row=i, column=2, value=fo)
        c.font = f(True, 11, "C00000" if i == 11 else "000000")
        c.alignment = Alignment(horizontal="center")
        for col in (1, 2):
            wr.cell(row=i, column=col).border = BORDE
        if et in COLORES_CLASIF:
            wr.cell(row=i, column=1).fill = PatternFill("solid", start_color=COLORES_CLASIF[et][0])

    # Mínimo por ambiente
    wr["D3"] = "Mínimo por ambiente"
    wr["D3"].font = f(True, 10, AZUL)
    for j, (amb, col) in enumerate((("Dev", "E"), ("QA", "F"), ("Prod", "G"))):
        c = wr.cell(row=4, column=5 + j, value=amb)
        c.font = f(True, 10, "FFFFFF"); c.fill = FILL_HEAD; c.alignment = Alignment(horizontal="center")
        for k, nivel in enumerate(("Obligatorio", "Recomendado", "Opcional")):
            wr.cell(row=5 + k, column=4, value=nivel).font = f()
            cc = wr.cell(row=5 + k, column=5 + j, value=f'=COUNTIF({rango(col)},"{nivel}")')
            cc.alignment = Alignment(horizontal="center"); cc.font = f(); cc.border = BORDE
    wr["D9"] = "Los niveles se pueden ajustar en vivo en las columnas Dev/QA/Prod del checklist."
    wr["D9"].font = f(size=8, italic=True, color=GRIS_TXT)

    # Por dominio
    fila = 14
    wr.cell(row=fila - 1, column=1, value="Por dominio").font = f(True, 11, AZUL)
    cabeza = ["Dominio", "Cumple", "Quick win", "Riesgo a validar", "Profundización", "Pendiente"]
    for j, h in enumerate(cabeza, start=1):
        c = wr.cell(row=fila, column=j, value=h)
        c.font = f(True, 9, "FFFFFF"); c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", wrap_text=True); c.border = BORDE
    wr.row_dimensions[fila].height = 28
    for i, (_, nom) in enumerate(DOMINIOS, start=fila + 1):
        wr.cell(row=i, column=1, value=nom).font = f(size=9)
        wr.cell(row=i, column=1).border = BORDE
        for j, et in enumerate(cabeza[1:], start=2):
            c = wr.cell(row=i, column=j,
                        value=f'=COUNTIFS({rango("B")},$A{i},{rango("L")},"{et}")')
            c.alignment = Alignment(horizontal="center"); c.font = f(size=9); c.border = BORDE
    for etiqueta, (fondo, texto) in COLORES_CLASIF.items():
        col = get_column_letter(cabeza.index(etiqueta) + 1)
        rr = f"{col}{fila + 1}:{col}{fila + len(DOMINIOS)}"
        wr.conditional_formatting.add(
            rr, FormulaRule(formula=[f"{col}{fila + 1}>0"],
                            fill=PatternFill("solid", start_color=fondo),
                            font=Font(name=FUENTE, bold=True, color=texto)))

    # Las tres listas — la salida formal del bloque
    fila = fila + len(DOMINIOS) + 3
    listas = [
        ("QUICK WINS — se ejecutan en semanas, con dueño", "O", 15, "Quick win"),
        ("RIESGOS A VALIDAR — nadie supo responder: se asigna quién averigua", "P", 15, "Riesgo a validar"),
        ("REQUIEREN SESIÓN DE PROFUNDIZACIÓN", "Q", 12, "Profundización"),
    ]
    for titulo, aux, filas, etiqueta in listas:
        c = wr.cell(row=fila, column=1, value=titulo)
        c.font = f(True, 11, COLORES_CLASIF[etiqueta][1])
        c.fill = PatternFill("solid", start_color=COLORES_CLASIF[etiqueta][0])
        wr.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=7)
        fila += 1
        for j, h in enumerate(("N.º", "ID", "Control", "", "", "Dueño", ""), start=1):
            if h:
                cc = wr.cell(row=fila, column=j, value=h)
                cc.font = f(True, 9, "FFFFFF"); cc.fill = FILL_HEAD
        wr.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=5)
        wr.merge_cells(start_row=fila, start_column=6, end_row=fila, end_column=7)
        fila += 1
        for k in range(1, filas + 1):
            buscar = f"MATCH({k},Checklist!${aux}${PRIMERA}:${aux}${ULTIMA},0)"
            wr.cell(row=fila, column=1, value=f'=IF(ISNUMBER({buscar}),{k},"")').font = f(size=9)
            wr.cell(row=fila, column=2,
                    value=f'=IFERROR(INDEX(Checklist!$A${PRIMERA}:$A${ULTIMA},{buscar}),"")').font = f(True, 9)
            c3 = wr.cell(row=fila, column=3,
                         value=f'=IFERROR(INDEX(Checklist!$C${PRIMERA}:$C${ULTIMA},{buscar}),"")')
            c3.font = f(size=9); c3.alignment = Alignment(wrap_text=True, vertical="top")
            wr.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=5)
            c6 = wr.cell(row=fila, column=6,
                         value=(f'=IFERROR(IF(INDEX(Checklist!$M${PRIMERA}:$M${ULTIMA},{buscar})="",'
                                f'"(sin dueño)",INDEX(Checklist!$M${PRIMERA}:$M${ULTIMA},{buscar})),"")'))
            c6.font = f(size=9)
            wr.merge_cells(start_row=fila, start_column=6, end_row=fila, end_column=7)
            wr.row_dimensions[fila].height = 26
            fila += 1
        fila += 1

    wr.conditional_formatting.add(
        f"F1:F{fila}", FormulaRule(formula=['F1="(sin dueño)"'],
                                   font=Font(name=FUENTE, bold=True, color="C00000")))
    wr.page_setup.orientation = "portrait"
    wr.sheet_properties.pageSetUpPr.fitToPage = True
    wr.page_setup.fitToWidth = 1
    wr.page_setup.fitToHeight = 0

    # ============================ LEEME ============================
    ws_leeme.column_dimensions["A"].width = 3
    ws_leeme.column_dimensions["B"].width = 26
    ws_leeme.column_dimensions["C"].width = 90
    ws_leeme["B2"] = "Checklist de seguridad — cómo se usa"
    ws_leeme["B2"].font = f(True, 16, AZUL)
    ws_leeme["B3"] = "Taller de arquitectura en OCI · módulo 2 · Seguridad cloud de extremo a extremo · la fecha del taller"
    ws_leeme["B3"].font = f(size=10, color=GRIS_TXT)

    filas = [
        ("En la sesión", None),
        ("1. Filtrar", "Hoja Checklist → filtro de la columna Núcleo = Sí. Son los controles que se revisan en vivo."),
        ("2. Preguntar", "Por cada control, la organización responde en «Estado la organización»: Sí · Parcial · No · No sé."),
        ("", "Si la respuesta es «Sí» inmediato, se pasa al siguiente. Solo se conversa en No, Parcial y No sé."),
        ("3. Ajustar mínimos", "Al cerrar cada dominio: «¿alguien no está de acuerdo con el mínimo propuesto para Dev, QA o Prod?»"),
        ("4. Dueños", "En quick wins y riesgos a validar se escribe un dueño. Nada sale de la sala sin dueño."),
        ("5. Resultado", "La hoja Resultado arma sola las tres listas que son la salida del bloque."),
        ("", None),
        ("Regla de clasificación", None),
        ("Sí", "→ Cumple"),
        ("No sé", "→ Riesgo a validar (nadie en la sala supo; se asigna quién averigua)"),
        ("No / Parcial + esfuerzo Bajo", "→ Quick win"),
        ("No / Parcial + esfuerzo Medio", "→ Profundización (necesita una sesión dedicada)"),
        ("", None),
        ("Qué se edita", None),
        ("Celdas amarillas", "Estado la organización, Dueño y Nota."),
        ("Dev / QA / Prod", "Nivel mínimo propuesto por ambiente. Se puede cambiar desde la lista desplegable."),
        ("Todo lo demás", "Se calcula. No escribir sobre Clasificación ni sobre la hoja Resultado."),
        ("", None),
        ("Ejemplo de fila llena", None),
    ]
    r = 5
    for a, b in filas:
        ca = ws_leeme.cell(row=r, column=2, value=a)
        if b is None and a:
            ca.font = f(True, 11, AZUL)
        else:
            ca.font = f(True, 10)
            cb = ws_leeme.cell(row=r, column=3, value=b)
            cb.font = f(size=10)
            cb.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    ejemplo_h = ["ID", "Control", "Estado", "Clasificación", "Dueño", "Nota"]
    ejemplo_v = ["I-01", "MFA obligatorio para todo acceso a la consola", "Parcial", "Quick win",
                 "Líder de infraestructura", "Solo administradores tienen MFA hoy"]
    for j, (h, v) in enumerate(zip(ejemplo_h, ejemplo_v)):
        cell_h = ws_leeme.cell(row=r, column=2, value=h)
        cell_h.font = f(True, 9, "FFFFFF"); cell_h.fill = FILL_HEAD; cell_h.border = BORDE
        cell_v = ws_leeme.cell(row=r, column=3, value=v)
        cell_v.font = f(size=9); cell_v.border = BORDE
        if h in ("Estado", "Dueño", "Nota"):
            cell_v.fill = FILL_INPUT
        if h == "Clasificación":
            fondo, texto = COLORES_CLASIF["Quick win"]
            cell_v.fill = PatternFill("solid", start_color=fondo)
            cell_v.font = Font(name=FUENTE, size=9, bold=True, color=texto)
        r += 1

    r += 1
    ws_leeme.cell(row=r, column=2, value="Fuente de los controles").font = f(True, 11, AZUL)
    r += 1
    nota = ("Selección propia para la conversación con la organización, alineada con prácticas de OCI y con el "
            "CIS OCI Foundations Benchmark. No es una auditoría ni una certificación. Los niveles mínimos "
            "por ambiente son una propuesta de partida para validar con el equipo de la organización. "
            "Se genera con checklist/generar_checklist.py; editar allí y regenerar.")
    c = ws_leeme.cell(row=r, column=3, value=nota)
    c.font = f(size=9, italic=True, color=GRIS_TXT)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws_leeme.row_dimensions[r].height = 52

    wb.active = 1  # abre en Checklist
    wb.save(ruta)


def construir_md(ruta):
    lineas = [
        "# Checklist de controles de seguridad",
        "",
        "Versión legible del checklist del módulo 2. **En la sesión se usa el Excel**",
        "(`Checklist_Seguridad.xlsx`), que clasifica cada respuesta automáticamente.",
        "Ambos se generan desde `generar_checklist.py`: editar allí, nunca a mano.",
        "",
        "**Niveles:** **O** = obligatorio · R = recomendado · Op = opcional · N/A = no aplica.",
        "**Núcleo** = se revisa en vivo en los 40 minutos; el resto queda como autoevaluación.",
        "",
        "**Regla de clasificación:** Sí → Cumple · No sé → Riesgo a validar ·",
        "No/Parcial con esfuerzo Bajo → Quick win · No/Parcial con esfuerzo Medio → Profundización.",
        "",
        f"Total: {len(CONTROLES)} controles · núcleo: {sum(1 for c in CONTROLES if c[8])}.",
        "",
    ]
    corto = {"O": "**O**", "R": "R", "Op": "Op", "NA": "N/A"}
    for letra, nombre in DOMINIOS:
        lineas += [f"## {nombre}", "",
                   "| ID | Control | Dev | QA | Prod | Esfuerzo | Núcleo |",
                   "|---|---|:-:|:-:|:-:|:-:|:-:|"]
        detalles = []
        for (cid, ctrl, porque, dev, qa, prod, esf, verif, nucleo) in CONTROLES:
            if cid[0] != letra:
                continue
            lineas.append(f"| **{cid}** | {ctrl} | {corto[dev]} | {corto[qa]} | {corto[prod]} | "
                          f"{esf} | {'●' if nucleo else ''} |")
            detalles.append(f"- **{cid}** — {porque} *Verificación:* {verif.replace('|', '&#124;')}")
        lineas += ["", *detalles, ""]
    ruta.write_text("\n".join(lineas), encoding="utf-8")


if __name__ == "__main__":
    ids = [c[0] for c in CONTROLES]
    assert len(ids) == len(set(ids)), "IDs duplicados en CONTROLES"
    construir_xlsx(AQUI / "Checklist_Seguridad.xlsx")
    construir_md(AQUI / "CHECKLIST.md")
    print(f"Generado: {len(CONTROLES)} controles, "
          f"{sum(1 for c in CONTROLES if c[8])} en el núcleo.")
