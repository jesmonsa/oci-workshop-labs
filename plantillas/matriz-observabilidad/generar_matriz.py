"""
Genera la salida formal del módulo 5 desde una única fuente de datos:

  - Matriz_Observabilidad.xlsx  -> se proyecta y se llena EN VIVO.
      · Hoja «Matriz»: señal -> umbral -> responsable -> acción.
      · Hoja «Resultado»: cobertura por categoría, brechas y agenda.
  - MATRIZ.md                -> la misma matriz en Markdown.

    python generar_matriz.py
Después abrir el Excel y guardarlo una vez para que calcule.

La mecánica que se explica en la sala:

    Una señal con umbral, responsable y acción es una alarma.
    Una señal que despierta a alguien SIN acción escrita es ruido.
    Y el ruido no es un problema menor: es lo que hace que dentro de seis meses
    nadie mire las alarmas que sí importan.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).parent

CATEGORIAS = ["Disponibilidad", "Latencia", "Errores", "Tráfico", "Saturación",
              "Capacidad", "Costo", "Seguridad", "Negocio", "Confiabilidad"]

# (id, categoría, señal, síntoma/causa, de dónde sale, sugerencia de umbral)
SENALES = [
    ("S-01", "Disponibilidad", "El servicio responde desde fuera (prueba sintética)",
     "Síntoma", "Monitoreo sintético cada minuto, desde fuera de la nube",
     "2 fallos seguidos"),
    ("S-02", "Latencia", "Latencia p95 de la API",
     "Síntoma", "Métricas del balanceador o instrumentación de la aplicación",
     "Acordar con el compromiso de servicio"),
    ("S-03", "Errores", "Tasa de respuestas 5xx",
     "Síntoma", "Métricas del balanceador", "> 1 % durante 5 min"),
    ("S-04", "Errores", "Fallos de integración con sistemas externos",
     "Síntoma", "Registros de la aplicación", "> N por hora"),
    ("S-05", "Tráfico", "Caída abrupta del volumen de transacciones",
     "Síntoma", "Base de datos o aplicación",
     "< 50 % de lo normal para esa hora"),
    ("S-06", "Saturación", "CPU de la capa de aplicación",
     "Causa", "Métricas de cómputo (la misma del autoescalamiento)",
     "> 45 % sostenido 3 min"),
    ("S-07", "Disponibilidad", "Servidores no saludables detrás del balanceador",
     "Síntoma", "Métricas del balanceador", "> 0 durante 2 min"),
    ("S-08", "Saturación", "Conexiones o memoria de la base de datos",
     "Causa", "Métricas de la base administrada", "> 80 % del máximo"),
    ("S-09", "Capacidad", "Almacenamiento de la base de datos",
     "Causa", "Métricas de la base administrada",
     "> 75 % (avisa con semanas, no con horas)"),
    ("S-10", "Capacidad", "El grupo de instancias lleva rato en su tamaño máximo",
     "Causa", "Métricas del grupo de instancias",
     "En el máximo más de 15 min"),
    ("S-11", "Negocio", "Transacciones detenidas en proceso",
     "Síntoma", "Consulta a la base (la del módulo 4)",
     "> N detenidas más de 4 h"),
    ("S-12", "Negocio", "Tiempo de proceso p95 por canal",
     "Síntoma", "Consulta a la base (la del módulo 4)",
     "Acordar con el compromiso de servicio"),
    ("S-13", "Costo", "Consumo del mes contra el presupuesto",
     "Causa", "Presupuestos y alertas (las del módulo 1)",
     "50 % · 75 % · 90 % · proyección"),
    ("S-14", "Seguridad", "Cambios en identidades y reglas de red",
     "Causa", "Eventos y notificaciones (las del módulo 2)",
     "Cualquier cambio, siempre"),
    ("S-15", "Seguridad", "Certificado TLS próximo a vencer",
     "Causa", "Servicio de certificados", "Faltan menos de 30 días"),
    ("S-16", "Confiabilidad", "Tiempo medio de recuperación del último mes",
     "Indicador", "Registro de incidentes",
     "No es alarma: se revisa cada mes"),
]

COLUMNAS = [
    ("ID", 7, None),
    ("Categoría", 15, None),
    ("Señal", 40, None),
    ("¿Síntoma o causa?", 13, None),
    ("De dónde sale", 34, None),
    ("Umbral sugerido", 26, None),
    ("Umbral acordado", 20, "input"),
    ("Severidad", 14, '"CRITICAL,WARNING,INFO,No es alarma"'),
    ("¿Existe hoy?", 12, '"Sí,Parcial,No,No sé"'),
    ("Responsable", 20, "input"),
    ("Acción o runbook", 28, "input"),
    ("¿Despierta a alguien?", 15, '"Sí 24x7,Solo en horario,No, solo tablero"'),
    ("Estado", 18, "formula"),
    ("Notas", 26, "input"),
]

FUENTE, AZUL, GRIS = "Arial", "1F3864", "595959"
FILL_HEAD = PatternFill("solid", start_color=AZUL)
FILL_INPUT = PatternFill("solid", start_color="FFF2CC")
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

COLORES_ESTADO = {
    "Lista": ("C6EFCE", "006100"),
    "Ruido": ("FFC7CE", "9C0006"),
    "Por implementar": ("BDD7EE", "1F3864"),
    "Incompleta": ("FFEB9C", "7F6000"),
}


def f(bold=False, size=10, color="000000", italic=False):
    return Font(name=FUENTE, bold=bold, size=size, color=color, italic=italic)


def construir_xlsx(ruta):
    wb = Workbook()
    ws_leeme = wb.active
    ws_leeme.title = "Leeme"
    wm = wb.create_sheet("Matriz")
    wr = wb.create_sheet("Resultado")

    n = len(SENALES)
    PRIMERA, ULTIMA = 5, 5 + n - 1 + 4          # cuatro filas libres para agregar
    ULTIMA_DATO = 5 + n - 1
    rango = lambda col: f"Matriz!${col}${PRIMERA}:${col}${ULTIMA}"

    # ============================ MATRIZ ============================
    wm["A1"] = "Matriz de observabilidad — Taller de arquitectura en OCI · módulo 5"
    wm["A1"].font = f(True, 14, AZUL)
    wm["A2"] = ("Señal → umbral → responsable → acción. La columna Estado se calcula sola: "
                "marca en rojo toda señal que despierte a alguien sin tener acción escrita.")
    wm["A2"].font = f(size=9, color=GRIS, italic=True)

    for j, (nombre, ancho, _) in enumerate(COLUMNAS, start=1):
        c = wm.cell(row=4, column=j, value=nombre)
        c.font = f(True, 10, "FFFFFF")
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        wm.column_dimensions[get_column_letter(j)].width = ancho
    wm.row_dimensions[4].height = 34

    for i in range(n + 4):
        r = PRIMERA + i
        datos = SENALES[i] if i < n else ("", "", "", "", "", "")
        for j, v in enumerate(datos, start=1):
            c = wm.cell(row=r, column=j, value=v)
            c.font = f(bold=(j == 1))
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="center" if j in (1, 4) else "left")
            c.border = BORDE
        for j, (_, _, tipo) in enumerate(COLUMNAS, start=1):
            if tipo in ("input",) or (tipo and tipo.startswith('"')):
                c = wm.cell(row=r, column=j)
                c.fill = FILL_INPUT
                c.border = BORDE
                c.font = f()
                c.alignment = Alignment(vertical="top", wrap_text=True,
                                        horizontal="center" if j in (8, 9, 12) else "left")
        # Estado
        c = wm.cell(row=r, column=13, value=(
            f'=IF($C{r}="","",'
            f'IF(I{r}="","Pendiente",'
            f'IF(AND(OR(L{r}="Sí 24x7",H{r}="CRITICAL"),OR(J{r}="",K{r}="")),"Ruido",'
            f'IF(OR(I{r}="No",I{r}="No sé"),"Por implementar",'
            f'IF(AND(G{r}<>"",J{r}<>"",K{r}<>""),"Lista","Incompleta")))))'))
        c.font = f(True)
        c.alignment = Alignment(horizontal="center", vertical="top")
        c.border = BORDE
        # Auxiliares para las listas del resultado
        for col, etiqueta in ((15, "Por implementar"), (16, "Ruido")):
            wm.cell(row=r, column=col,
                    value=f'=IF($M{r}="{etiqueta}",COUNTIF($M${PRIMERA}:$M{r},"{etiqueta}"),"")'
                    ).font = f(size=8, color="A6A6A6")
        wm.row_dimensions[r].height = 42

    for col in ("O", "P"):
        wm.column_dimensions[col].hidden = True

    for j, (_, _, tipo) in enumerate(COLUMNAS, start=1):
        if tipo and tipo.startswith('"'):
            dv = DataValidation(type="list", formula1=tipo, allow_blank=True)
            wm.add_data_validation(dv)
            letra = get_column_letter(j)
            dv.add(f"{letra}{PRIMERA}:{letra}{ULTIMA}")

    for etiqueta, (fondo, texto) in COLORES_ESTADO.items():
        wm.conditional_formatting.add(
            f"M{PRIMERA}:M{ULTIMA}",
            FormulaRule(formula=[f'$M{PRIMERA}="{etiqueta}"'],
                        fill=PatternFill("solid", start_color=fondo),
                        font=Font(name=FUENTE, bold=True, color=texto)))

    wm.auto_filter.ref = f"A4:N{ULTIMA}"
    wm.freeze_panes = "D5"
    wm.page_setup.orientation = "landscape"
    wm.sheet_properties.pageSetUpPr.fitToPage = True
    wm.page_setup.fitToWidth = 1
    wm.page_setup.fitToHeight = 0
    wm.print_title_rows = "4:4"

    # ============================ RESULTADO ============================
    wr["A1"] = "Modelo de observabilidad mínimo viable — se arma solo"
    wr["A1"].font = f(True, 14, AZUL)
    for col, ancho in zip("ABCDEFG", [24, 13, 13, 13, 13, 13, 40]):
        wr.column_dimensions[col].width = ancho

    wr["A3"] = "Indicador"
    wr["B3"] = "Valor"
    for c in (wr["A3"], wr["B3"]):
        c.font = f(True, 10, "FFFFFF")
        c.fill = FILL_HEAD
        c.border = BORDE

    kpis = [
        ("Señales en la matriz", f'=COUNTA({rango("C")})'),
        ("Listas (umbral, dueño y acción)", f'=COUNTIF({rango("M")},"Lista")'),
        ("Por implementar", f'=COUNTIF({rango("M")},"Por implementar")'),
        ("Incompletas", f'=COUNTIF({rango("M")},"Incompleta")'),
        ("RUIDO: despiertan sin acción", f'=COUNTIF({rango("M")},"Ruido")'),
        ("Sin evaluar", f'=COUNTIF({rango("M")},"Pendiente")'),
        ("Señales que despiertan 24x7",
         f'=COUNTIF({rango("L")},"Sí 24x7")'),
    ]
    for i, (etq, fo) in enumerate(kpis, start=4):
        c1 = wr.cell(row=i, column=1, value=etq)
        c1.font = f(bold=(etq.startswith("RUIDO")))
        c1.border = BORDE
        c2 = wr.cell(row=i, column=2, value=fo)
        c2.font = f(True, 11, "C00000" if etq.startswith("RUIDO") else "000000")
        c2.alignment = Alignment(horizontal="center")
        c2.border = BORDE
    wr.cell(row=11, column=3,
            value="Si este número no es cero, la agenda empieza aquí: "
                  "escribir la acción o bajar la severidad.").font = f(size=9, italic=True, color="C00000")
    wr.cell(row=10, column=3,
            value="Más de 6 ó 7 señales que despiertan de madrugada es una promesa "
                  "que ningún equipo pequeño sostiene.").font = f(size=9, italic=True, color=GRIS)

    # Cobertura por categoría
    fila = 14
    wr.cell(row=fila - 1, column=1, value="Cobertura por categoría").font = f(True, 11, AZUL)
    cabeza = ["Categoría", "Señales", "Listas", "Faltan", "¿Cubierta?"]
    for j, h in enumerate(cabeza, start=1):
        c = wr.cell(row=fila, column=j, value=h)
        c.font = f(True, 9, "FFFFFF")
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BORDE
    for i, cat in enumerate(CATEGORIAS, start=fila + 1):
        wr.cell(row=i, column=1, value=cat).font = f(size=9)
        wr.cell(row=i, column=1).border = BORDE
        wr.cell(row=i, column=2, value=f'=COUNTIF({rango("B")},$A{i})')
        wr.cell(row=i, column=3, value=f'=COUNTIFS({rango("B")},$A{i},{rango("M")},"Lista")')
        wr.cell(row=i, column=4, value=f'=B{i}-C{i}')
        wr.cell(row=i, column=5, value=f'=IF(C{i}>0,"Sí","NO")')
        for j in range(2, 6):
            c = wr.cell(row=i, column=j)
            c.alignment = Alignment(horizontal="center")
            c.font = f(size=9)
            c.border = BORDE
    wr.conditional_formatting.add(
        f"A{fila+1}:E{fila+len(CATEGORIAS)}",
        FormulaRule(formula=[f'$E{fila+1}="NO"'],
                    fill=PatternFill("solid", start_color="FFC7CE")))
    wr.cell(row=fila + len(CATEGORIAS) + 1, column=1,
            value="Una categoría sin ninguna señal lista es un punto ciego. "
                  "Disponibilidad y Errores son las dos que no se pueden dejar vacías.").font = \
        f(size=9, italic=True, color=GRIS)

    # Listas automáticas
    fila = fila + len(CATEGORIAS) + 4
    for titulo, aux, filas_n, etiqueta in (
            ("PRIMERO: señales que despiertan sin acción escrita", "P", 6, "Ruido"),
            ("DESPUÉS: señales por implementar", "O", 14, "Por implementar")):
        c = wr.cell(row=fila, column=1, value=titulo)
        c.font = f(True, 11, COLORES_ESTADO[etiqueta][1])
        c.fill = PatternFill("solid", start_color=COLORES_ESTADO[etiqueta][0])
        wr.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=7)
        fila += 1
        for j, h in enumerate(("N.º", "ID", "Señal", "", "", "Responsable", "Para cuándo"), start=1):
            if h:
                cc = wr.cell(row=fila, column=j, value=h)
                cc.font = f(True, 9, "FFFFFF")
                cc.fill = FILL_HEAD
        wr.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=5)
        fila += 1
        for k in range(1, filas_n + 1):
            buscar = f"MATCH({k},Matriz!${aux}${PRIMERA}:${aux}${ULTIMA},0)"
            wr.cell(row=fila, column=1, value=f'=IF(ISNUMBER({buscar}),{k},"")').font = f(size=9)
            wr.cell(row=fila, column=2,
                    value=f'=IFERROR(INDEX(Matriz!$A${PRIMERA}:$A${ULTIMA},{buscar}),"")'
                    ).font = f(True, 9)
            c3 = wr.cell(row=fila, column=3,
                         value=f'=IFERROR(INDEX(Matriz!$C${PRIMERA}:$C${ULTIMA},{buscar}),"")')
            c3.font = f(size=9)
            c3.alignment = Alignment(wrap_text=True, vertical="top")
            wr.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=5)
            c6 = wr.cell(row=fila, column=6,
                         value=(f'=IFERROR(IF(INDEX(Matriz!$J${PRIMERA}:$J${ULTIMA},{buscar})="",'
                                f'"(sin dueño)",INDEX(Matriz!$J${PRIMERA}:$J${ULTIMA},{buscar})),"")'))
            c6.font = f(size=9)
            c7 = wr.cell(row=fila, column=7)
            c7.fill = FILL_INPUT
            c7.border = BORDE
            wr.row_dimensions[fila].height = 24
            fila += 1
        fila += 1

    # Agenda
    wr.cell(row=fila, column=1, value="Agenda de implementación").font = f(True, 11, AZUL)
    fila += 1
    for j, h in enumerate(("Momento", "Qué se pone en marcha", "", "", "", "Responsable", "Listo el"), start=1):
        if h:
            c = wr.cell(row=fila, column=j, value=h)
            c.font = f(True, 9, "FFFFFF")
            c.fill = FILL_HEAD
    wr.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=5)
    fila += 1
    for momento in ("Esta semana", "Próximas 2 semanas", "Dentro de 90 días"):
        wr.cell(row=fila, column=1, value=momento).font = f(True, 9)
        wr.cell(row=fila, column=1).border = BORDE
        for j in (2, 6, 7):
            c = wr.cell(row=fila, column=j)
            c.fill = FILL_INPUT
            c.border = BORDE
        wr.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=5)
        wr.row_dimensions[fila].height = 30
        fila += 1

    wr.page_setup.orientation = "portrait"
    wr.sheet_properties.pageSetUpPr.fitToPage = True
    wr.page_setup.fitToWidth = 1
    wr.page_setup.fitToHeight = 0

    # ============================ LEEME ============================
    ws_leeme.column_dimensions["A"].width = 3
    ws_leeme.column_dimensions["B"].width = 30
    ws_leeme.column_dimensions["C"].width = 92
    ws_leeme["B2"] = "Cómo se usa este archivo"
    ws_leeme["B2"].font = f(True, 16, AZUL)
    ws_leeme["B3"] = "Taller de arquitectura en OCI · módulo 5 · Monitoreo, observabilidad y confiabilidad"
    ws_leeme["B3"].font = f(size=10, color=GRIS)

    bloques = [
        ("En la sesión", None),
        ("Minuto 17", "Hoja «Matriz». Se recorre señal por señal para UN servicio crítico. "
                      "Las columnas amarillas se llenan en voz alta."),
        ("Por cada señal", "Cuatro preguntas: ¿qué umbral? · ¿quién responde? · ¿qué hace? · "
                           "¿esto despierta a alguien de madrugada?"),
        ("Minuto 31", "Hoja «Resultado». La cobertura y las dos listas ya están armadas."),
        ("Minuto 34", "La agenda: esta semana, dos semanas, 90 días. Con nombre y fecha."),
        ("", None),
        ("La regla del bloque", None),
        ("Una alarma sin acción es ruido",
         "Si una señal despierta a alguien y no tiene acción escrita, la matriz la marca en "
         "rojo. Hay dos formas de arreglarlo: escribir la acción, o bajarle la severidad. "
         "Las dos son respuestas válidas; dejarla como está, no."),
        ("Síntoma antes que causa",
         "Los síntomas los siente el usuario (el servicio no responde, va lento, da error). "
         "Las causas explican por qué (CPU, memoria, disco). Si hay que elegir por dónde "
         "empezar, se empieza por los síntomas: son los que corresponden a un impacto real."),
        ("«No sé» vale", "Marca dónde no hay visibilidad. Es el resultado del ejercicio."),
        ("", None),
        ("Qué se edita", None),
        ("Celdas amarillas", "Umbral acordado, severidad, existe hoy, responsable, acción, "
                             "si despierta, notas y las fechas de la agenda. Lo demás se calcula."),
        ("", None),
        ("Después de la sesión", None),
        ("Entrega", "Este archivo va en la memoria del bloque, dentro de las 48 horas. "
                    "La primera lista —las que despiertan sin acción— se resuelve en una semana: "
                    "es escribir, no construir."),
    ]
    r = 5
    for a, b in bloques:
        ca = ws_leeme.cell(row=r, column=2, value=a)
        if b is None and a:
            ca.font = f(True, 11, AZUL)
        else:
            ca.font = f(True, 10)
            cb = ws_leeme.cell(row=r, column=3, value=b)
            cb.font = f(size=10)
            cb.alignment = Alignment(wrap_text=True, vertical="top")
            ws_leeme.row_dimensions[r].height = 42
        r += 1

    wb.active = 1
    wb.save(ruta)


def construir_md():
    lineas = [
        "# Matriz de observabilidad — señales candidatas",
        "",
        "Catálogo de partida para el módulo 5. **En la sesión se usa el Excel**",
        "(`matriz/Matriz_Observabilidad.xlsx`), que calcula la cobertura y",
        "marca solo las señales que despiertan a alguien sin tener acción escrita.",
        "",
        "Las señales vienen de los bloques anteriores: la CPU del autoescalamiento (módulo 1),",
        "los cambios críticos de seguridad (módulo 2), las transacciones detenidas y los",
        "tiempos por canal (módulo 4), y el consumo contra presupuesto (módulo 1).",
        "",
        "**Estado (se calcula solo):**",
        "",
        "| Estado | Cuándo |",
        "|---|---|",
        "| **Lista** | Existe, con umbral acordado, responsable y acción |",
        "| **Ruido** | Despierta a alguien (24x7 o CRITICAL) sin responsable o sin acción |",
        "| **Por implementar** | No existe todavía |",
        "| **Incompleta** | Existe pero le falta umbral acordado, responsable o acción |",
        "",
        f"Total: {len(SENALES)} señales candidatas en {len(CATEGORIAS)} categorías.",
        "",
        "| ID | Categoría | Señal | ¿Síntoma o causa? | De dónde sale | Umbral sugerido |",
        "|---|---|---|:-:|---|---|",
    ]
    for sid, cat, senal, tipo, fuente, umbral in SENALES:
        lineas.append(f"| **{sid}** | {cat} | {senal} | {tipo} | {fuente} | {umbral} |")
    lineas += [
        "",
        "## Lo que se llena en vivo",
        "",
        "Umbral acordado · Severidad · ¿Existe hoy? · **Responsable** · **Acción o runbook** ·",
        "¿Despierta a alguien? · Notas.",
        "",
        "> Las dos columnas en negrita son las que convierten una señal en una alarma.",
        "> Sin ellas hay una gráfica bonita y un teléfono que suena de madrugada sin que",
        "> nadie sepa qué hacer.",
        "",
    ]
    (AQUI / "MATRIZ.md").write_text("\n".join(lineas), encoding="utf-8")


if __name__ == "__main__":
    ids = [s[0] for s in SENALES]
    assert len(ids) == len(set(ids)), "IDs duplicados"
    for s in SENALES:
        assert s[1] in CATEGORIAS, f"categoría desconocida: {s[1]}"
    construir_xlsx(AQUI / "Matriz_Observabilidad.xlsx")
    construir_md()
    print(f"Generado: {len(SENALES)} señales en {len(CATEGORIAS)} categorías.")
